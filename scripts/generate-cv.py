"""Generate bilingual ATS CVs from content/profile.json.

Dependencies: python-docx, reportlab. Run: python scripts/generate-cv.py
The generated files contain no source PDFs, phone number or private address.
"""
from pathlib import Path
import json
from html import escape
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.pagesizes import A4

ROOT = Path(__file__).resolve().parents[1]
PROFILE = json.loads((ROOT / 'content/profile.json').read_text())
LABELS = {
 'pt': {'summary':'Resumo','experience':'Experiência','continued':'Experiência anterior','education':'Formação','skills':'Habilidades','languages':'Idiomas'},
 'en': {'summary':'Summary','experience':'Experience','continued':'Earlier experience','education':'Education','skills':'Skills','languages':'Languages'},
}

# A plain single-column layout is shared by both export formats.
# The explicit page break keeps the first three roles together and avoids
# leaving a role heading or an isolated bullet at a page boundary.
def blocks(lang):
 p=PROFILE[lang]; labels=LABELS[lang]
 yield 'title', PROFILE['name']
 yield 'subtitle', p['title']
 location='Natal, RN, Brasil' if lang=='pt' else 'Natal, Brazil'
 yield 'contact', f"{location} | {PROFILE['email']}"
 yield 'contact', f"{PROFILE['github'].replace('https://','')} | {PROFILE['linkedin'].replace('https://www.','').rstrip('/')}"
 yield 'heading', labels['summary']
 yield 'body', p['summary']
 yield 'heading', labels['experience']
 for i,e in enumerate(p['experience']):
  if i==3:
   yield 'break',''
   yield 'heading', labels['continued']
  yield 'role', f"{e['company']} | {e['role']}"
  yield 'meta', f"{e['period']} | {e['location']}"
  for b in e['bullets']: yield 'bullet',b
 yield 'heading', labels['education']
 for e in p['education']:
  yield 'role', e['degree']
  yield 'body', f"{e['institution']} | {e['period']}"
 yield 'heading',labels['skills']
 for s in p['skills']: yield 'skill', (s['category'], ', '.join(s['items']))
 yield 'heading',labels['languages']
 yield 'body',' | '.join(p['languages'])

def make_docx(lang):
 d=Document(); sec=d.sections[0]
 sec.page_width=Cm(21); sec.page_height=Cm(29.7)
 sec.top_margin=Cm(1.65); sec.bottom_margin=Cm(1.65)
 sec.left_margin=Cm(1.8); sec.right_margin=Cm(1.8)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
  st=d.styles[name]
  st.font.name='Liberation Sans'; st.font.size=Pt(10.5); st.font.color.rgb=RGBColor(0,0,0)
  st.font.italic=False
  for border in list(st.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
  for spacing in list(st.element.iter(qn('w:spacing'))):
   if spacing.getparent().tag==qn('w:rPr'): spacing.getparent().remove(spacing)
  st.paragraph_format.space_after=Pt(4)
  st.paragraph_format.line_spacing=1.08
  st.paragraph_format.widow_control=True
 d.styles['Title'].font.size=Pt(23)
 d.styles['Title'].font.bold=True
 d.styles['Subtitle'].font.size=Pt(12)
 d.styles['Subtitle'].font.bold=True
 d.styles['Heading 1'].font.size=Pt(11)
 d.styles['Heading 1'].font.bold=True
 d.styles['Heading 1'].paragraph_format.space_before=Pt(11)
 d.styles['Heading 1'].paragraph_format.space_after=Pt(5)
 d.styles['Heading 1'].paragraph_format.keep_with_next=True
 d.styles['Heading 2'].font.bold=True
 d.styles['Heading 2'].paragraph_format.space_before=Pt(7)
 d.styles['Heading 2'].paragraph_format.space_after=Pt(2)
 d.styles['Heading 2'].paragraph_format.keep_with_next=True
 d.core_properties.title=PROFILE['name']+' - '+PROFILE[lang]['title']
 d.core_properties.author=PROFILE['name']
 d.core_properties.subject='Curriculum vitae'
 for kind,text in blocks(lang):
  if kind=='break': d.add_page_break();continue
  style={'title':'Title','subtitle':'Subtitle','heading':'Heading 1','role':'Heading 2','bullet':'Normal'}.get(kind,'Normal')
  para=d.add_paragraph(style=style)
  if kind=='skill':
   para.add_run(text[0]+': ').bold=True;para.add_run(text[1])
  else: para.add_run(('• ' if kind=='bullet' else '')+text)
  if kind=='contact':
   para.paragraph_format.space_after=Pt(2)
  if kind=='meta':
   para.paragraph_format.keep_with_next=True
   para.paragraph_format.space_after=Pt(3)
  if kind=='bullet':
   para.paragraph_format.left_indent=Cm(0.35)
   para.paragraph_format.first_line_indent=Cm(-0.35)
  for run in para.runs:
   el=run._element.get_or_add_rPr(); language=OxmlElement('w:lang');language.set(qn('w:val'),'pt-BR' if lang=='pt' else 'en-US');el.append(language)
 out=ROOT/'cvs'/f'Lucas-Gabriel-Ferreira-CV-{lang}.docx'
 out.parent.mkdir(exist_ok=True);d.save(out)

def make_pdf(lang):
 base=dict(fontName='Helvetica',fontSize=10.5,leading=13,spaceAfter=4,textColor='#000000')
 styles={
 'body':ParagraphStyle('body',**base),
 'title':ParagraphStyle('title',fontName='Helvetica-Bold',fontSize=23,leading=26,spaceAfter=5),
 'subtitle':ParagraphStyle('subtitle',fontName='Helvetica-Bold',fontSize=12,leading=15,spaceAfter=5),
 'contact':ParagraphStyle('contact',fontName='Helvetica',fontSize=10.5,leading=13,spaceAfter=2),
 'heading':ParagraphStyle('heading',fontName='Helvetica-Bold',fontSize=11,leading=14,spaceBefore=11,spaceAfter=5,keepWithNext=True),
 'role':ParagraphStyle('role',fontName='Helvetica-Bold',fontSize=10.5,leading=13,spaceBefore=7,spaceAfter=2,keepWithNext=True),
 'meta':ParagraphStyle('meta',fontName='Helvetica',fontSize=10.5,leading=13,spaceAfter=3,keepWithNext=True),
 'bullet':ParagraphStyle('bullet',fontName='Helvetica',fontSize=10.5,leading=13,spaceAfter=4,leftIndent=10,firstLineIndent=-7),
 }
 flow=[]
 for kind,text in blocks(lang):
  if kind=='break':flow.append(PageBreak());continue
  if kind=='skill': html='<b>'+escape(text[0])+':</b> '+escape(text[1]);kind='body'
  else:html=escape(text).replace('’',"'")
  if kind=='bullet':html='&#8226; '+html
  flow.append(Paragraph(html,styles[kind]))
 out=ROOT/'public/cv'/f'Lucas-Gabriel-Ferreira-CV-{lang}.pdf'
 out.parent.mkdir(exist_ok=True,parents=True)
 SimpleDocTemplate(str(out),pagesize=A4,leftMargin=51,rightMargin=51,topMargin=47,bottomMargin=47,title=PROFILE['name']+' - '+PROFILE[lang]['title'],author=PROFILE['name']).build(flow)

if __name__=='__main__':
 for lang in ('pt','en'):
  make_docx(lang);make_pdf(lang)
  print(f'Generated {lang}: PDF and DOCX')
